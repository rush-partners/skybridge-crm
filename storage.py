"""Almacenamiento en disco de los archivos adjuntos de una importación
(despacho, invoices, packing list, documentación de carga, facturas de pago)."""
import os
import re
import subprocess
import tempfile
import threading
import uuid
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

DATA_DIR = Path(os.environ.get("DATA_DIR", Path(__file__).parent))
UPLOADS_DIR = DATA_DIR / "uploads"
ARCHIVOS_CLIENTES_DIR = DATA_DIR / "archivos_clientes"

# soffice headless no tolera bien dos conversiones en paralelo sobre el mismo
# perfil de usuario (colisiona el lock del perfil) — un lock de proceso
# alcanza porque Streamlit corre todo en un solo proceso Python.
_conversion_lock = threading.Lock()


def _forzar_ajuste_pagina(ruta_xlsx: Path):
    """Fuerza 'área de impresión = todo el rango usado' + 'ajustar a 1
    página' en cada hoja, reescribiendo el .xlsx en el lugar.
    Sin esto, LibreOffice decide sus propios saltos de página según el
    layout original del archivo, y con archivos reales (con imágenes
    flotantes ancladas a celdas, como fotos de producto en un invoice) el
    filtro SinglePageSheets del export termina RECORTANDO las primeras
    filas en vez de escalar todo el contenido — se vio con un invoice real
    que perdía el encabezado del fabricante. Fijar el área de impresión
    explícitamente al rango completo evita esa ambigüedad."""
    wb = openpyxl.load_workbook(ruta_xlsx)
    for ws in wb.worksheets:
        if ws.max_row < 1 or ws.max_column < 1:
            continue
        ultima_col = get_column_letter(ws.max_column)
        ws.print_area = f"A1:{ultima_col}{ws.max_row}"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        if ws.sheet_properties.pageSetUpPr is None:
            ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        else:
            ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(ruta_xlsx)


def convertir_a_pdf(contenido: bytes, nombre_original: str) -> bytes | None:
    """Convierte un Excel (.xlsx/.xls) a PDF vía LibreOffice headless.
    Devuelve None si la conversión falla (LibreOffice no instalado, archivo
    corrupto, timeout, etc.) en vez de lanzar — el llamador decide cómo
    avisarle al usuario."""
    ext = (Path(nombre_original).suffix or ".xlsx").lower()
    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        origen = tmp_path / f"input{ext}"
        origen.write_bytes(contenido)

        ajustado = False
        if ext == ".xlsx":
            # openpyxl no lee el formato binario viejo .xls — para esos cae
            # al filtro SinglePageSheets de más abajo.
            try:
                _forzar_ajuste_pagina(origen)
                ajustado = True
            except Exception:
                pass

        # El filtro SinglePageSheets del export escala a 1 sola página, pero
        # ignora/pisa el área de impresión ya fijada arriba y vuelve a
        # recortar las primeras filas (mismo bug) — por eso solo se usa como
        # fallback cuando NO se pudo pre-ajustar el archivo (.xls, o algún
        # .xlsx que falló al reescribirse).
        destino = "pdf"
        if not ajustado:
            destino = 'pdf:calc_pdf_Export:{"SinglePageSheets":{"type":"boolean","value":true}}'
        try:
            with _conversion_lock:
                resultado = subprocess.run(
                    ["soffice", "--headless", "--convert-to", destino, "--outdir", str(tmp_path), str(origen)],
                    capture_output=True, timeout=60,
                )
        except (subprocess.TimeoutExpired, FileNotFoundError):
            return None
        if resultado.returncode != 0:
            return None
        salida = tmp_path / "input.pdf"
        return salida.read_bytes() if salida.exists() else None


def _slug(texto: str) -> str:
    return re.sub(r"[^a-zA-Z0-9._-]+", "_", texto).strip("_") or "archivo"


def guardar_archivo(importacion_id: int, categoria: str, nombre_original: str, contenido: bytes) -> str:
    carpeta = UPLOADS_DIR / f"importacion_{importacion_id}" / _slug(categoria)
    carpeta.mkdir(parents=True, exist_ok=True)
    nombre_seguro = f"{uuid.uuid4().hex[:8]}_{_slug(nombre_original)}"
    ruta = carpeta / nombre_seguro
    ruta.write_bytes(contenido)
    return str(ruta)


def guardar_archivo_cliente(cliente_id: int, nombre_original: str, contenido: bytes) -> str:
    carpeta = ARCHIVOS_CLIENTES_DIR / str(cliente_id)
    carpeta.mkdir(parents=True, exist_ok=True)
    nombre_seguro = f"{uuid.uuid4().hex[:8]}_{_slug(nombre_original)}"
    ruta = carpeta / nombre_seguro
    ruta.write_bytes(contenido)
    return str(ruta)


def eliminar_archivo(ruta: str):
    p = Path(ruta)
    if p.exists():
        p.unlink()


def leer_archivo(ruta: str) -> bytes:
    p = Path(ruta)
    return p.read_bytes() if p.exists() else b""
